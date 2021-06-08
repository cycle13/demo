from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt,Cm


def tabless(numi,numj,document,list,table_style):
    table = document.add_table(rows=numi, cols=numj)
    for i in range(len(table.rows)):
        # table.rows[i].height = Cm(2)
        for j in range(len(table.columns)):
            table.cell(i, j).text=list[i*len(table.columns)+j]
            # table.cell(i, j).width = Cm(10)


    for row in table.rows:
          for cell in row.cells:
                paragraphs = cell.paragraphs
                for paragraph in paragraphs:
                      paragraph_format = paragraph.paragraph_format
                      paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                      paragraph_format.space_before = Pt(5)  # 上行间距
                      paragraph_format.space_after = Pt(5)  # 下行间距
                      for run in paragraph.runs:
                            font = run.font
                            font.size= Pt(10.5)

    table.style = table_style  #设置表格样式